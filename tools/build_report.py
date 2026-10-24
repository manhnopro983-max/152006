from pathlib import Path
import re
try:
    from docx import Document  # type: ignore
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.shared import Inches, Pt, RGBColor  # type: ignore
except ImportError:
    Document = None  # type: ignore
    WD_ALIGN_PARAGRAPH = None  # type: ignore
    Inches = Pt = RGBColor = None  # type: ignore


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs' / 'BAO_CAO.md'
TARGET = ROOT / 'docs' / 'BAO_CAO_EXPENSE_TRACKER.docx'

if Document is None:
    print("Thông báo: Cần cài đặt python-docx nếu muốn tạo lại file Word: pip install python-docx")
    import sys
    sys.exit(0)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(5)
for name, size, color in [('Title', 26, '245A86'), ('Heading 1', 17, '245A86'), ('Heading 2', 13, '245A86'), ('Heading 3', 11, '2E6F5E')]:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)

lines = SOURCE.read_text(encoding='utf-8').splitlines()
cover = doc.add_paragraph()
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
cover.paragraph_format.space_before = Pt(100)
r = cover.add_run('BÁO CÁO MÔN HỌC')
r.bold = True
r.font.name = 'Arial'
r.font.size = Pt(27)
r.font.color.rgb = RGBColor.from_string('245A86')
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_before = Pt(18)
r = subtitle.add_run('XÂY DỰNG HỆ THỐNG QUẢN LÝ CHI TIÊU CÁ NHÂN')
r.bold = True
r.font.size = Pt(18)

cover_info = doc.add_paragraph()
cover_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
cover_info.paragraph_format.space_before = Pt(55)
cover_info.add_run('Sinh viên: [Họ và tên]\nMSSV: [Mã số sinh viên]\nLớp: [Tên lớp]\nGiảng viên: [Họ tên giảng viên]\nNgày nộp: [Ngày/tháng/năm]').font.size = Pt(12)
doc.add_page_break()
doc.add_heading('Mục lục', level=1)
for item in ['1. Giới thiệu', '2. Kiến trúc hệ thống', '3. Ứng dụng web và cơ sở dữ liệu', '4. Docker Compose, Nginx và pgAdmin', '5. Prometheus và Grafana', '6. Loki, Promtail và LogQL', '7. Hardening bảo mật', '8. Kết quả triển khai và demo', '9. Kết luận', '10. Tài liệu tham khảo']:
    doc.add_paragraph(item, style='List Number')
doc.add_paragraph('Các phần chính bắt đầu trên trang riêng; số trang sẽ được cập nhật trong Word sau khi điền thông tin cá nhân.')

def clean(text):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    text = re.sub(r'`([^`]+)`', r'\1', text)
    text = text.replace('**', '').replace('*', '')
    return text.strip()

def add_table(block):
    rows = []
    for line in block:
        cells = [clean(c.strip()) for c in line.strip().strip('|').split('|')]
        if all(re.fullmatch(r':?-{3,}:?', c or '-') for c in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    width = max(map(len, rows))
    table = doc.add_table(rows=0, cols=width)
    table.style = 'Light Shading Accent 1'
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for j in range(width):
            cells[j].text = row[j] if j < len(row) else ''
            for p in cells[j].paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for run in p.runs:
                    run.font.size = Pt(8.5)
                    if index == 0:
                        run.bold = True

# Cover details and the Markdown table of contents have already been rendered
# above; start body conversion at the first numbered report section.
i = next(index for index, line in enumerate(lines) if line.startswith('## 1. '))
in_code = False
code_lines = []
while i < len(lines):
    line = lines[i]
    if line.startswith('## '):
        title = line[3:].strip()
        if re.match(r'\d+\.', title):
            heading = doc.add_heading(clean(title), level=1)
            heading.paragraph_format.page_break_before = True
        elif title in ('Mục lục', 'Đề tài: Xây dựng Hệ thống Quản lý Chi tiêu Cá nhân'):
            pass
        else:
            doc.add_heading(clean(title), level=1)
        i += 1
        continue
    if line.startswith('### '):
        doc.add_heading(clean(line[4:]), level=2)
        i += 1
        continue
    if line.startswith('#### '):
        doc.add_heading(clean(line[5:]), level=3)
        i += 1
        continue
    if line.startswith('```'):
        in_code = not in_code
        if not in_code and code_lines:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run('\n'.join(code_lines))
            run.font.name = 'Consolas'
            run.font.size = Pt(8)
            code_lines = []
        i += 1
        continue
    if in_code:
        code_lines.append(line)
        i += 1
        continue
    if line.startswith('|'):
        block = []
        while i < len(lines) and lines[i].startswith('|'):
            block.append(lines[i]); i += 1
        add_table(block)
        continue
    image_match = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', line.strip())
    if image_match:
        image_path = (SOURCE.parent / image_match.group(2)).resolve()
        if image_path.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Two Grafana views share one page; keep them compact enough that
            # the second screenshot caption cannot be orphaned onto a blank page.
            width = Inches(4.8) if image_path.name.startswith('grafana-') else Inches(6.35)
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(image_path), width=width)
            caption = doc.add_paragraph(image_match.group(1))
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.keep_together = True
            for run in caption.runs:
                run.italic = True
                run.font.size = Pt(8)
        i += 1
        continue
    if line.startswith('[*Chèn Screenshot'):
        doc.add_paragraph('Ảnh minh chứng này được chụp khi demo trực tiếp; xem checklist ảnh trong docs/DEMO_CHECKLIST.md.')
        i += 1
        continue
    if line.startswith('---') or not line.strip():
        i += 1
        continue
    if line.startswith('> '):
        p = doc.add_paragraph(clean(line[2:]))
        p.paragraph_format.left_indent = Inches(0.25)
        for run in p.runs:
            run.italic = True
        i += 1
        continue
    if re.match(r'^- \[[ xX]\]', line):
        text = re.sub(r'^- \[[ xX]\]\s*', '', line)
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(('☑ ' if '- [x]' in line.lower() else '☐ ') + clean(text))
        i += 1
        continue
    if line.startswith('- '):
        doc.add_paragraph(clean(line[2:]), style='List Bullet')
        i += 1
        continue
    if re.match(r'^\d+\.\s', line):
        doc.add_paragraph(clean(re.sub(r'^\d+\.\s', '', line)), style='List Number')
        i += 1
        continue
    doc.add_paragraph(clean(line))
    i += 1

for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Hệ thống Quản lý Chi tiêu | Báo cáo môn học').font.size = Pt(8)

doc.save(TARGET)
print(TARGET)
